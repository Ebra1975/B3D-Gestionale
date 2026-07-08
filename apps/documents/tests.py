import shutil
import tempfile
import zipfile
from io import BytesIO
from decimal import Decimal

from django.test import TestCase, override_settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from django.utils import timezone
from docx import Document as DocxDocument

from apps.customers.models import Customer
from apps.documents.forms import DocumentTemplateForm
from apps.documents.models import DocumentProfile, DocumentTemplate, GeneratedDocument
from apps.documents.services import (
    build_document_export_checks,
    generate_consulting_estimate_docx,
    generate_internal_estimate_docx,
    get_or_create_default_consulting_template,
)
from apps.estimates.models import CostItem, Estimate, EstimateConfiguration
from apps.inventory.models import Material


def make_template_upload(filename="template.docx", paragraphs=None):
    buffer = BytesIO()
    document = DocxDocument()
    for paragraph in paragraphs or ["Template {{ preventivo.numero }} per {{ cliente.nome }}"]:
        document.add_paragraph(paragraph)
    document.save(buffer)
    buffer.seek(0)
    return SimpleUploadedFile(
        filename,
        buffer.read(),
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


def read_docx_word_xml(docx_path):
    with zipfile.ZipFile(docx_path) as docx:
        return "\n".join(
            docx.read(name).decode("utf-8")
            for name in docx.namelist()
            if name.startswith("word/") and name.endswith(".xml")
        )


class DocumentExportTests(TestCase):
    def setUp(self):
        self.media_root = tempfile.mkdtemp(prefix="b3dlab_test_media_")
        self.override = override_settings(MEDIA_ROOT=self.media_root)
        self.override.enable()

    def tearDown(self):
        self.override.disable()
        shutil.rmtree(self.media_root, ignore_errors=True)

    def test_export_checks_block_empty_estimate(self):
        estimate = Estimate.objects.create(
            number="B3D-DOC-EMPTY",
            customer=Customer.objects.create(name="Cliente Documenti"),
            subject="",
            date=timezone.localdate(),
        )

        checks = build_document_export_checks(estimate)

        self.assertFalse(checks["can_generate"])
        self.assertIn("Completare l'oggetto del preventivo.", checks["missing_items"])
        self.assertIn("Inserire una descrizione della richiesta cliente.", checks["missing_items"])
        self.assertIn("Scegliere o aggiungere una configurazione tecnica.", checks["missing_items"])

    def test_internal_document_is_generated_from_same_cost_items(self):
        DocumentProfile.objects.create(
            name="Profilo test",
            company_name="B3D Lab Test",
            fiscal_note="Nota fiscale da validare con commercialista.",
            active=True,
        )
        material = Material.objects.create(name="PETG", cost_per_unit=Decimal("20.00"))
        estimate = Estimate.objects.create(
            number="B3D-DOC-001",
            customer=Customer.objects.create(name="Cliente PDF", email="cliente@example.com"),
            subject="Supporto tecnico",
            description="Realizzazione e validazione di un componente tecnico.",
            date=timezone.localdate(),
            valid_until=timezone.localdate(),
            quantity=2,
        )
        configuration = EstimateConfiguration.objects.create(
            estimate=estimate,
            name="Configurazione scelta",
            description="Soluzione FDM in PETG.",
            material=material,
            quantity=2,
            is_selected=True,
        )
        CostItem.objects.create(
            configuration=configuration,
            category=CostItem.Category.MATERIAL,
            description="Materiale PETG",
            quantity=Decimal("1.00"),
            unit="kg",
            unit_cost=Decimal("20.00"),
        )
        CostItem.objects.create(
            configuration=configuration,
            category=CostItem.Category.MARGIN,
            description="Margine commerciale",
            quantity=Decimal("1.00"),
            unit="voce",
            unit_cost=Decimal("10.00"),
        )

        document = generate_internal_estimate_docx(estimate)

        self.assertEqual(document.document_type, GeneratedDocument.DocumentType.INTERNAL)
        self.assertEqual(document.version, 1)
        self.assertTrue(document.docx_file.name.endswith("_interno_v1.docx"))
        self.assertEqual(document.template.template_version, "v2")
        self.assertTrue(document.template.template_file.name.endswith("preventivo_interno_base_v2.docx"))
        self.assertEqual(configuration.total, Decimal("30.00"))

    def test_consulting_document_uses_sprint_11_template_without_internal_costs(self):
        DocumentProfile.objects.create(
            name="Profilo test cliente",
            company_name="B3D Lab Test",
            fiscal_note="Nota fiscale da validare con commercialista.",
            active=True,
        )
        material = Material.objects.create(name="PETG", cost_per_unit=Decimal("20.00"))
        estimate = Estimate.objects.create(
            number="B3D-DOC-CLI",
            customer=Customer.objects.create(name="Cliente PDF", email="cliente@example.com"),
            subject="Supporto tecnico",
            description="Realizzazione e validazione di un componente tecnico.",
            date=timezone.localdate(),
            valid_until=timezone.localdate(),
            quantity=2,
        )
        configuration = EstimateConfiguration.objects.create(
            estimate=estimate,
            name="Configurazione scelta",
            description="Soluzione FDM in PETG.",
            material=material,
            quantity=2,
            is_selected=True,
        )
        CostItem.objects.create(
            configuration=configuration,
            category=CostItem.Category.MATERIAL,
            description="Materiale PETG",
            quantity=Decimal("1.00"),
            unit="kg",
            unit_cost=Decimal("20.00"),
        )
        CostItem.objects.create(
            configuration=configuration,
            category=CostItem.Category.MARGIN,
            description="Margine commerciale",
            quantity=Decimal("1.00"),
            unit="voce",
            unit_cost=Decimal("10.00"),
        )

        document = generate_consulting_estimate_docx(estimate)

        self.assertEqual(document.document_type, GeneratedDocument.DocumentType.CONSULTING)
        self.assertEqual(document.template.template_version, "v3")
        self.assertTrue(document.template.template_file.name.endswith("preventivo_consulenza_base_v3.docx"))
        self.assertTrue(document.docx_file.name.endswith("_consulenza_v1.docx"))
        with zipfile.ZipFile(document.docx_file.path) as docx:
            document_xml = docx.read("word/document.xml").decode("utf-8")
        self.assertNotIn("Costo interno", document_xml)
        self.assertNotIn("Margine percentuale", document_xml)
        self.assertNotIn("Voci di costo", document_xml)

    def test_realistic_estimate_generates_client_and_internal_documents(self):
        DocumentProfile.objects.create(
            name="Profilo prova reale",
            company_name="B3D Lab",
            subtitle="Consulenza tecnica e manifattura additiva",
            email="info@b3dlab.example",
            phone="3330000000",
            standard_consulting_terms="Validita 15 giorni. Tempi da confermare dopo accettazione.",
            fiscal_note="Dicitura fiscale da validare con commercialista.",
            internal_footer_note="Documento interno: non inviare al cliente.",
            active=True,
        )
        material = Material.objects.create(
            name="PETG CF",
            brand="Bambu Lab",
            color="Nero",
            cost_per_unit=Decimal("48.00"),
        )
        customer = Customer.objects.create(
            name="Officina Rossi",
            contact_person="Marco Rossi",
            email="marco.rossi@example.com",
            tax_code="RSSMRC80A01H501U",
        )
        estimate = Estimate.objects.create(
            number="B3D-2026-REAL",
            customer=customer,
            subject="Supporto tecnico per staffa attrezzaggio",
            description=(
                "Analisi del componente, piccola ottimizzazione geometrica, "
                "stampa di validazione e consegna di due prototipi funzionali."
            ),
            date=timezone.localdate(),
            valid_until=timezone.localdate(),
            quantity=2,
            general_terms="Pagamento e tempi da confermare dopo validazione tecnica.",
        )
        configuration = EstimateConfiguration.objects.create(
            estimate=estimate,
            name="Prototipo funzionale PETG CF",
            description="Soluzione FDM rinforzata per verifica ingombri e montaggio.",
            material=material,
            process="FDM",
            treatment="Rimozione supporti e controllo dimensionale base",
            expected_duration="3 giorni lavorativi",
            operating_mode="Consulenza con realizzazione prototipo",
            public_notes="Il prototipo serve per validazione tecnica prima di eventuale serie.",
            internal_notes="Controllare orientamento pezzo prima della stampa definitiva.",
            quantity=2,
            is_selected=True,
        )
        CostItem.objects.create(
            configuration=configuration,
            category=CostItem.Category.SETUP,
            description="Analisi file e preparazione stampa",
            quantity=Decimal("1.00"),
            unit="voce",
            unit_cost=Decimal("35.00"),
        )
        CostItem.objects.create(
            configuration=configuration,
            category=CostItem.Category.MATERIAL,
            description="Materiale PETG CF",
            quantity=Decimal("0.35"),
            unit="kg",
            unit_cost=Decimal("48.00"),
        )
        CostItem.objects.create(
            configuration=configuration,
            category=CostItem.Category.MACHINE_TIME,
            description="Tempo macchina stimato",
            quantity=Decimal("6.00"),
            unit="h",
            unit_cost=Decimal("6.00"),
        )
        CostItem.objects.create(
            configuration=configuration,
            category=CostItem.Category.MARGIN,
            description="Margine commerciale",
            quantity=Decimal("1.00"),
            unit="voce",
            unit_cost=Decimal("42.00"),
        )

        client_document = generate_consulting_estimate_docx(estimate)
        internal_document = generate_internal_estimate_docx(estimate)

        self.assertEqual(client_document.version, 1)
        self.assertEqual(internal_document.version, 1)
        self.assertTrue(client_document.docx_file.name.endswith("_consulenza_v1.docx"))
        self.assertTrue(internal_document.docx_file.name.endswith("_interno_v1.docx"))
        self.assertEqual(configuration.total, Decimal("129.80"))

        client_xml = read_docx_word_xml(client_document.docx_file.path)
        internal_xml = read_docx_word_xml(internal_document.docx_file.path)

        self.assertIn("Officina Rossi", client_xml)
        self.assertIn("Supporto tecnico per staffa attrezzaggio", client_xml)
        self.assertIn("129.80 EUR", client_xml)
        self.assertNotIn("Costo interno", client_xml)
        self.assertIn("Costo interno", internal_xml)
        self.assertIn("Analisi file e preparazione stampa", internal_xml)
        self.assertIn("Documento interno: non inviare al cliente.", internal_xml)


class DocumentProfileViewTests(TestCase):
    def test_document_profile_can_be_updated_from_interface(self):
        profile = DocumentProfile.get_active()

        response = self.client.post(
            reverse("documents:profile_update"),
            {
                "name": "Profilo documenti B3D",
                "company_name": "B3D Lab aggiornata",
                "subtitle": "Consulenza tecnica e manifattura additiva",
                "address": "Via di prova 1",
                "email": "info@b3dlab.example",
                "phone": "3330000000",
                "website": "https://b3dlab.example",
                "tax_code": "IT00000000000",
                "standard_consulting_terms": "Condizioni standard aggiornate.",
                "fiscal_note": "Nota fiscale aggiornata da validare con commercialista.",
                "internal_footer_note": "Documento interno aggiornato.",
            },
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(response["Location"], reverse("documents:list"))
        profile.refresh_from_db()
        self.assertEqual(profile.company_name, "B3D Lab aggiornata")
        self.assertEqual(profile.standard_consulting_terms, "Condizioni standard aggiornate.")
        self.assertEqual(profile.fiscal_note, "Nota fiscale aggiornata da validare con commercialista.")


class DocumentTemplateViewTests(TestCase):
    def setUp(self):
        self.media_root = tempfile.mkdtemp(prefix="b3dlab_test_media_")
        self.override = override_settings(MEDIA_ROOT=self.media_root)
        self.override.enable()

    def tearDown(self):
        self.override.disable()
        shutil.rmtree(self.media_root, ignore_errors=True)

    def test_template_can_be_uploaded_from_interface(self):
        response = self.client.post(
            reverse("documents:template_create"),
            {
                "name": "Template consulenza personalizzato",
                "document_type": DocumentTemplate.DocumentType.CONSULTING_ESTIMATE,
                "profile": "consulting",
                "template_version": "cliente-v1",
                "active": "on",
                "notes": "Template caricato da interfaccia.",
                "template_file": make_template_upload("template_cliente.docx"),
            },
        )

        self.assertEqual(response.status_code, 302)
        template = DocumentTemplate.objects.get(name="Template consulenza personalizzato")
        self.assertTrue(template.active)
        self.assertTrue(template.template_file.name.endswith(".docx"))

    def test_template_upload_rejects_non_docx_files(self):
        form = DocumentTemplateForm(
            {
                "name": "Template errato",
                "document_type": DocumentTemplate.DocumentType.CONSULTING_ESTIMATE,
                "profile": "consulting",
                "template_version": "v1",
                "active": "on",
            },
            {"template_file": SimpleUploadedFile("template.txt", b"test", content_type="text/plain")},
        )

        self.assertFalse(form.is_valid())
        self.assertIn("Il template deve essere un file .docx.", form.errors["template_file"])
        self.assertFalse(DocumentTemplate.objects.filter(name="Template errato").exists())

    def test_template_upload_rejects_renamed_invalid_docx(self):
        form = DocumentTemplateForm(
            {
                "name": "Template rinominato",
                "document_type": DocumentTemplate.DocumentType.CONSULTING_ESTIMATE,
                "profile": "consulting",
                "template_version": "v1",
                "active": "on",
            },
            {
                "template_file": SimpleUploadedFile(
                    "template.docx",
                    b"non e davvero un file Word",
                    content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )

        self.assertFalse(form.is_valid())
        self.assertIn("Il file non e un DOCX valido", form.errors["template_file"][0])
        self.assertFalse(DocumentTemplate.objects.filter(name="Template rinominato").exists())

    def test_template_upload_rejects_unknown_placeholders_for_supported_documents(self):
        form = DocumentTemplateForm(
            {
                "name": "Template con segnaposto errato",
                "document_type": DocumentTemplate.DocumentType.CONSULTING_ESTIMATE,
                "profile": "consulting",
                "template_version": "v1",
                "active": "on",
            },
            {"template_file": make_template_upload("template.docx", ["Cliente {{ client.nome }}"])},
        )

        self.assertFalse(form.is_valid())
        self.assertIn("client", form.errors["template_file"][0])
        self.assertFalse(DocumentTemplate.objects.filter(name="Template con segnaposto errato").exists())

    def test_activating_template_deactivates_same_document_type(self):
        first = DocumentTemplate.objects.create(
            name="Template consulenza A",
            document_type=DocumentTemplate.DocumentType.CONSULTING_ESTIMATE,
            profile="consulting",
            template_file="templates/consulting_estimate/a.docx",
            template_version="a",
            active=True,
        )
        second = DocumentTemplate.objects.create(
            name="Template consulenza B",
            document_type=DocumentTemplate.DocumentType.CONSULTING_ESTIMATE,
            profile="consulting",
            template_file="templates/consulting_estimate/b.docx",
            template_version="b",
            active=False,
        )

        response = self.client.post(reverse("documents:template_activate", args=[second.pk]))

        self.assertEqual(response.status_code, 302)
        first.refresh_from_db()
        second.refresh_from_db()
        self.assertFalse(first.active)
        self.assertTrue(second.active)

    def test_custom_active_template_is_used_before_generated_default(self):
        default = DocumentTemplate.objects.create(
            name="Preventivo consulenza base",
            document_type=DocumentTemplate.DocumentType.CONSULTING_ESTIMATE,
            profile="consulting",
            template_file="templates/consulting_estimate/preventivo_consulenza_base_v3.docx",
            template_version="v3",
            active=True,
        )
        custom = DocumentTemplate.objects.create(
            name="Template consulenza cliente",
            document_type=DocumentTemplate.DocumentType.CONSULTING_ESTIMATE,
            profile="consulting",
            template_file="templates/consulting_estimate/custom.docx",
            template_version="cliente-v1",
            active=True,
        )

        selected = get_or_create_default_consulting_template()

        self.assertEqual(selected, custom)
        default.refresh_from_db()
        self.assertEqual(default.template_version, "v3")
